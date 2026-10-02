import io
import sqlite3
from hashlib import sha256
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from httpx import Response
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from backend.api.main import create_app
from backend.config import Settings


def client(tmp_path: Path) -> TestClient:
    return TestClient(create_app(Settings(data_dir=tmp_path)))


def upload(
    api: TestClient,
    files: list[tuple[str, bytes]],
    *,
    document_ids: str | None = None,
    chunk_size: int = 4,
    chunk_overlap: int = 1,
) -> Response:
    data = {"name": "知识文档", "chunk_size": str(chunk_size), "chunk_overlap": str(chunk_overlap)}
    if document_ids is not None:
        data["document_ids"] = document_ids
    return api.post(
        "/api/v1/document-collections",
        data=data,
        files=[("files", (name, io.BytesIO(content))) for name, content in files],
    )


def docx_file() -> bytes:
    document = Document()
    document.add_paragraph("First paragraph")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Left"
    table.cell(0, 1).text = "Right"
    document.add_paragraph("Last paragraph")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def pdf_file(*, text: str | None = None, protected: bool = False) -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    if text is not None:
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
        )
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 72 250 Td ({text}) Tj ET".encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    if protected:
        writer.encrypt("secret")
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def test_upload_preserves_bytes_ids_checksum_and_ordered_chunks(tmp_path: Path) -> None:
    api = client(tmp_path)
    content = b"\xef\xbb\xbfabcdef"
    result = upload(
        api,
        [("guide.md", content), ("guide.txt", "甲乙丙丁".encode())],
        document_ids='["rag-guide", null]',
    )
    assert result.status_code == 201
    detail = result.json()
    assert detail["chunk_size"] == 4
    assert detail["source_kind"] == "original_files"
    assert all(doc["has_original_file"] for doc in detail["documents"])
    assert [chunk["position"] for chunk in detail["chunks"]] == [0, 1, 2]
    assert detail["chunk_overlap"] == 1
    assert [doc["document_id"] for doc in detail["documents"]] == ["rag-guide", "guide-1"]
    assert detail["documents"][0]["checksum"] == sha256(content).hexdigest()
    assert [chunk["text"] for chunk in detail["documents"][0]["chunks"]] == ["abcd", "def"]
    assert [chunk["text"] for chunk in detail["documents"][1]["chunks"]] == ["甲乙丙丁"]
    with sqlite3.connect(tmp_path / "rageva.sqlite3") as connection:
        stored = connection.execute(
            "SELECT content FROM source_documents WHERE document_id = ?", ("rag-guide",)
        ).fetchone()[0]
    assert stored == content
    assert api.get(f"/api/v1/document-collections/{detail['id']}").json() == detail
    assert api.get("/api/v1/document-collections").json()[0]["document_count"] == 2


def test_duplicate_names_get_distinct_generated_ids(tmp_path: Path) -> None:
    result = upload(
        client(tmp_path),
        [("same.txt", b"one"), ("same.md", b"two"), ("manual.txt", b"three")],
        document_ids='[null, null, "same-1"]',
    )
    assert result.status_code == 201
    assert [doc["document_id"] for doc in result.json()["documents"]] == [
        "same-2",
        "same-3",
        "same-1",
    ]


def test_invalid_file_rejects_entire_collection_with_file_issue(tmp_path: Path) -> None:
    api = client(tmp_path)
    result = upload(api, [("good.txt", b"valid"), ("bad.md", b"\xff")])
    assert result.status_code == 422
    assert result.json()["issues"][0]["file_index"] == 1
    assert result.json()["issues"][0]["code"] == "invalid_encoding"
    assert api.get("/api/v1/document-collections").json() == []
    with sqlite3.connect(tmp_path / "rageva.sqlite3") as connection:
        assert connection.execute("SELECT COUNT(*) FROM source_documents").fetchone()[0] == 0


def test_ids_empty_files_formats_and_chunk_settings_are_rejected(tmp_path: Path) -> None:
    api = client(tmp_path)
    cases = [
        ([("a.txt", b"a"), ("b.txt", b"b")], '["same", "same"]', 4, 1, "duplicate_document_id"),
        ([("a.txt", b"a")], '[""]', 4, 1, "empty_document_id"),
        ([("a.txt", b"")], None, 4, 1, "empty_file"),
        ([("a.bin", b"content")], None, 4, 1, "unsupported_format"),
        ([("a.txt", b"content")], None, 4, 4, "invalid_chunk_overlap"),
        ([("a.txt", b"content")], None, 0, 0, "invalid_chunk_size"),
    ]
    for files, ids, size, overlap, code in cases:
        result = upload(api, files, document_ids=ids, chunk_size=size, chunk_overlap=overlap)
        assert result.status_code == 422
        assert code in {issue["code"] for issue in result.json()["issues"]}
    assert api.get("/api/v1/document-collections").json() == []


def test_snapshot_does_not_change_after_new_upload(tmp_path: Path) -> None:
    api = client(tmp_path)
    first = upload(api, [("a.txt", b"abcdef")]).json()
    second = upload(api, [("a.txt", b"new text")], chunk_size=6, chunk_overlap=0).json()
    assert first["id"] != second["id"]
    assert api.get(f"/api/v1/document-collections/{first['id']}").json() == first


def test_docx_and_text_pdf_extract_in_order_but_keep_original_bytes(tmp_path: Path) -> None:
    api = client(tmp_path)
    originals = [("report.docx", docx_file()), ("brief.pdf", pdf_file(text="PDF evidence"))]
    response = upload(api, originals, chunk_size=1000, chunk_overlap=0)
    assert response.status_code == 201
    detail = response.json()
    documents = detail["documents"]
    assert [document["document_id"] for document in documents] == ["report-1", "brief-1"]
    assert documents[0]["chunks"][0]["text"] == ("First paragraph\nLeft\tRight\nLast paragraph")
    assert documents[1]["chunks"][0]["text"] == "PDF evidence"
    with sqlite3.connect(tmp_path / "rageva.sqlite3") as connection:
        stored = connection.execute(
            "SELECT filename, content, checksum FROM source_documents ORDER BY position"
        ).fetchall()
    for index, (filename, content, checksum) in enumerate(stored):
        assert filename == originals[index][0]
        assert content == originals[index][1]
        assert checksum == sha256(originals[index][1]).hexdigest()
    assert api.get(f"/api/v1/document-collections/{detail['id']}").json() == detail


def test_invalid_binary_files_reject_the_entire_collection(tmp_path: Path) -> None:
    api = client(tmp_path)
    bad_files = [
        ("broken.docx", b"not a DOCX", "invalid_docx"),
        ("broken.pdf", b"not a PDF", "invalid_pdf"),
        ("scan.pdf", pdf_file(), "no_extractable_text"),
        ("locked.pdf", pdf_file(text="secret", protected=True), "protected_pdf"),
    ]
    for filename, content, code in bad_files:
        response = upload(api, [("valid.txt", b"valid"), (filename, content)])
        assert response.status_code == 422
        issues = response.json()["issues"]
        assert len(issues) == 1
        assert issues[0]["file_index"] == 1
        assert issues[0]["filename"] == filename
        assert issues[0]["field"] == "file"
        assert issues[0]["code"] == code
        assert issues[0]["message"]
    assert api.get("/api/v1/document-collections").json() == []
    with sqlite3.connect(tmp_path / "rageva.sqlite3") as connection:
        assert connection.execute("SELECT COUNT(*) FROM source_documents").fetchone()[0] == 0
